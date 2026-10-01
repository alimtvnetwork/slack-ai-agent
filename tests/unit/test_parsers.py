from __future__ import annotations

import io

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from slack_agent.files.parser import (
    parse_code_bytes,
    parse_csv_bytes,
    parse_docx_bytes,
    parse_file_bytes,
    parse_html_bytes,
    parse_pdf_bytes,
)


def _generate_test_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.drawString(100, 750, "Hello Slack Agent PDF!")
    c.showPage()
    c.save()
    return buffer.getvalue()


def _generate_test_docx_bytes() -> bytes:
    buffer = io.BytesIO()
    doc = Document()
    doc.add_heading("Test Document Title", level=1)
    doc.add_paragraph("First paragraph content.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col1"
    table.cell(0, 1).text = "Col2"
    table.cell(1, 0).text = "Val1"
    table.cell(1, 1).text = "Val2"
    doc.save(buffer)
    return buffer.getvalue()


def test_parse_pdf_bytes_success() -> None:
    pdf_bytes = _generate_test_pdf_bytes()
    result = parse_pdf_bytes(pdf_bytes)

    assert result.is_success is True
    assert "Hello Slack Agent PDF!" in result.value()


def test_parse_docx_bytes_success() -> None:
    docx_bytes = _generate_test_docx_bytes()
    result = parse_docx_bytes(docx_bytes)

    assert result.is_success is True
    content = result.value()
    assert "Test Document Title" in content
    assert "First paragraph content." in content
    assert "Col1 | Col2" in content
    assert "Val1 | Val2" in content


def test_parse_csv_bytes_success() -> None:
    csv_bytes = b"id,name,role\n1,Alice,Engineer\n2,Bob,Architect"
    result = parse_csv_bytes(csv_bytes)

    assert result.is_success is True
    content = result.value()
    assert "id | name | role" in content
    assert "1 | Alice | Engineer" in content


def test_parse_html_bytes_success() -> None:
    html_bytes = b"""<!DOCTYPE html>
    <html>
    <head><title>Course Overview - Telehandler</title><style>.hidden{display:none;}</style></head>
    <body>
    <script>console.log('strip me');</script>
    <h1>KI Training & Assessing</h1>
    <p>Nationally recognised training for RIIHAN309F Conduct operations.</p>
    <footer>Copyright 2026</footer>
    </body>
    </html>"""
    result = parse_html_bytes(html_bytes)

    assert result.is_success is True
    content = result.value()
    assert "Course Overview - Telehandler" in content
    assert "KI Training & Assessing" in content
    assert "RIIHAN309F" in content
    assert "console.log" not in content
    assert "Copyright 2026" not in content


def test_parse_code_bytes_success() -> None:
    code_bytes = b"def hello() -> str:\n    return 'world'"
    result = parse_code_bytes(code_bytes)

    assert result.is_success is True
    assert "def hello() -> str:" in result.value()


def test_parse_file_bytes_dispatch() -> None:
    pdf_bytes = _generate_test_pdf_bytes()
    res_pdf = parse_file_bytes(pdf_bytes, "report.pdf", "application/pdf")
    assert res_pdf.is_success is True

    docx_bytes = _generate_test_docx_bytes()
    res_docx = parse_file_bytes(docx_bytes, "spec.docx", "application/vnd.openxmlformats")
    assert res_docx.is_success is True

    csv_bytes = b"k,v\n1,2"
    res_csv = parse_file_bytes(csv_bytes, "data.csv", "text/csv")
    assert res_csv.is_success is True

    html_bytes = b"<html><head><title>Test</title></head><body><p>Article</p></body></html>"
    res_html = parse_file_bytes(html_bytes, "page.html", "text/html")
    assert res_html.is_success is True
    assert "Article" in res_html.value()

    code_bytes = b"print('hi')"
    res_py = parse_file_bytes(code_bytes, "main.py", "text/x-python")
    assert res_py.is_success is True


def test_parse_pdf_invalid_bytes_fails() -> None:
    result = parse_pdf_bytes(b"not-a-valid-pdf-content")
    assert result.has_error is True
