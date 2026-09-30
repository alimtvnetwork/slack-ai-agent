from __future__ import annotations

import io

from docx import Document

from slack_agent.core.errors import AppError, ErrorCategoryType
from slack_agent.core.result import Result


def generate_docx_bytes(title: str, content: str) -> Result[bytes]:
    """Render Word DOCX document bytes using python-docx."""
    try:
        buffer = io.BytesIO()
        doc = Document()

        # Add Title Heading
        doc.add_heading(title, level=1)

        # Add paragraphs
        for block in content.split("\n\n"):
            clean_block = block.strip()
            if clean_block:
                doc.add_paragraph(clean_block)

        doc.save(buffer)
        return Result.ok(buffer.getvalue())
    except Exception as exc:
        return Result.fail(
            AppError.wrap(
                exc,
                code="E5004",
                message="python-docx failed to build DOCX document",
                category=ErrorCategoryType.FileSystem,
                context={"Title": title},
            )
        )
