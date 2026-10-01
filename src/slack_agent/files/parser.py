from __future__ import annotations

import csv
import io
import re

import openpyxl
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from slack_agent.core.errors import (
    ERR_FILE_PARSE_FAILED,
    AppError,
    ErrorCategoryType,
)
from slack_agent.core.result import Result

_HTML_TAGS_TO_REMOVE = (
    "script",
    "style",
    "nav",
    "footer",
    "header",
    "noscript",
    "aside",
    "svg",
    "iframe",
)


def parse_pdf_bytes(file_bytes: bytes) -> Result[str]:
    """Extract plain text from PDF byte content using pypdf."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text: list[str] = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(f"--- Page {index + 1} ---\n{text.strip()}")

        extracted = "\n\n".join(pages_text)
        return Result.ok(extracted)
    except Exception as exc:
        return Result.fail(
            AppError.wrap(
                exc,
                code=ERR_FILE_PARSE_FAILED,
                message="Failed to parse PDF document bytes",
                category=ErrorCategoryType.FileSystem,
            )
        )


def parse_docx_bytes(file_bytes: bytes) -> Result[str]:
    """Extract text and tables from Word DOCX byte content."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        elements: list[str] = []

        # Extract paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                elements.append(text)

        # Extract table rows
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    elements.append(" | ".join(row_cells))

        extracted = "\n\n".join(elements)
        return Result.ok(extracted)
    except Exception as exc:
        return Result.fail(
            AppError.wrap(
                exc,
                code=ERR_FILE_PARSE_FAILED,
                message="Failed to parse DOCX document bytes",
                category=ErrorCategoryType.FileSystem,
            )
        )


def parse_csv_bytes(file_bytes: bytes) -> Result[str]:
    """Format CSV byte content as structured rows."""
    try:
        text = file_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [" | ".join(row) for row in reader if row]
        return Result.ok("\n".join(rows))
    except Exception as exc:
        return Result.fail(
            AppError.wrap(
                exc,
                code=ERR_FILE_PARSE_FAILED,
                message="Failed to parse CSV bytes",
                category=ErrorCategoryType.FileSystem,
            )
        )


def parse_code_bytes(file_bytes: bytes) -> Result[str]:
    """Decode text or code file bytes using UTF-8."""
    try:
        text = file_bytes.decode("utf-8", errors="replace")
        return Result.ok(text)
    except Exception as exc:
        return Result.fail(
            AppError.wrap(
                exc,
                code=ERR_FILE_PARSE_FAILED,
                message="Failed to decode text file bytes",
                category=ErrorCategoryType.FileSystem,
            )
        )


def parse_excel_bytes(file_bytes: bytes) -> Result[str]:
    """Extract worksheets and tabular rows from Excel .xlsx byte content."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sections: list[str] = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            sheet_rows: list[str] = [f"--- Sheet: {sheet_name} ---"]
            for row in sheet.iter_rows(values_only=True):
                row_values = [
                    str(val).strip() for val in row if val is not None and str(val).strip()
                ]
                if row_values:
                    sheet_rows.append(" | ".join(row_values))

            if len(sheet_rows) > 1:
                sections.append("\n".join(sheet_rows))

        extracted = "\n\n".join(sections)
        return Result.ok(extracted)
    except Exception as exc:
        return Result.fail(
            AppError.wrap(
                exc,
                code=ERR_FILE_PARSE_FAILED,
                message="Failed to parse Excel document bytes",
                category=ErrorCategoryType.FileSystem,
            )
        )


def parse_html_bytes(file_bytes: bytes) -> Result[str]:
    """Extract readable text from HTML markup byte content."""
    try:
        html_text = file_bytes.decode("utf-8", errors="replace")
        soup = BeautifulSoup(html_text, "html.parser")
        for tag in soup(_HTML_TAGS_TO_REMOVE):
            tag.decompose()

        page_title = soup.title.string.strip() if soup.title and soup.title.string else ""
        body_text = "\n".join(soup.stripped_strings)
        clean_text = re.sub(r"\n{3,}", "\n\n", body_text)
        full_text = f"# {page_title}\n\n{clean_text}" if page_title else clean_text
        return Result.ok(full_text)
    except Exception as exc:
        return Result.fail(
            AppError.wrap(
                exc,
                code=ERR_FILE_PARSE_FAILED,
                message="Failed to parse HTML document bytes",
                category=ErrorCategoryType.FileSystem,
            )
        )


def parse_file_bytes(file_bytes: bytes, file_name: str, file_type: str) -> Result[str]:
    """Dispatch file byte parsing based on detected extension or file type."""
    lower_name = file_name.lower()
    lower_type = file_type.lower()

    if lower_name.endswith(".pdf") or lower_type == "pdf":
        return parse_pdf_bytes(file_bytes)

    if lower_name.endswith((".html", ".htm")) or "html" in lower_type:
        return parse_html_bytes(file_bytes)

    if lower_name.endswith(".docx") or "word" in lower_type:
        return parse_docx_bytes(file_bytes)

    if (
        lower_name.endswith((".xlsx", ".xls"))
        or "spreadsheet" in lower_type
        or "excel" in lower_type
    ):
        return parse_excel_bytes(file_bytes)

    if lower_name.endswith(".csv") or lower_type == "csv":
        return parse_csv_bytes(file_bytes)

    return parse_code_bytes(file_bytes)
